import docx

doc = docx.Document(r'C:\Users\Srinithy\OneDrive\Desktop\minetrack_backup.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"Total sections: {len(doc.sections)}")

print("\n--- SCANNING ALL DRAWINGS AND FIGURES ---")
for i, p in enumerate(doc.paragraphs):
    drawings = p._p.xpath('.//*[local-name()="drawing" or local-name()="pict"]')
    blips = p._p.xpath('.//*[local-name()="blip" or local-name()="imagedata"]')
    txt = p.text.strip()
    
    if drawings or blips or 'figure' in txt.lower() or 'fig' in txt.lower():
        prev_p = doc.paragraphs[i-1].text.strip() if i > 0 else ""
        next_p = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ""
        print(f"P{i:3d}: drawings={len(drawings)}, blips={len(blips)}")
        print(f"     text: {repr(txt)}")
        if drawings:
            print(f"     prev: {repr(prev_p[:70])}")
            print(f"     next: {repr(next_p[:70])}")

print("\n--- TABLES SUMMARY ---")
for i, t in enumerate(doc.tables):
    rows = len(t.rows)
    cols = len(t.columns)
    first_cells = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]
    print(f"Table {i:2d}: {rows} rows x {cols} cols | {first_cells[:3]}")
