import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import zipfile

DOC_PATH = r'C:\Users\Srinithy\OneDrive\Desktop\MineTrack_Project_Report_Formatted.docx'

doc = docx.Document(DOC_PATH)

def setup_page_number_footer(section, is_roman=False, start_num=1):
    sectPr = section._sectPr
    
    # Configure pgNumType
    pg_num_type = sectPr.find(qn('w:pgNumType'))
    if pg_num_type is None:
        pg_num_type = OxmlElement('w:pgNumType')
        sectPr.append(pg_num_type)
    
    fmt = 'lowerRoman' if is_roman else 'decimal'
    pg_num_type.set(qn('w:fmt'), fmt)
    pg_num_type.set(qn('w:start'), str(start_num))
    
    # Configure footer
    footer = section.footer
    footer.is_linked_to_previous = False
    
    # Clear existing paragraphs
    for p in list(footer.paragraphs):
        footer._element.remove(p._p)
        
    p = footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    run = p.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    
    # Insert field for PAGE
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    run._r.append(fldSimple)

def clear_section_footers(section):
    sectPr = section._sectPr
    # Remove any footer references or clear footer
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in list(footer.paragraphs):
        footer._element.remove(p._p)
    p = footer.add_paragraph()
    p.text = ""

# Section 0: Cover & Certificates -> NO page numbers
clear_section_footers(doc.sections[0])

# Section 1: Preliminary Pages -> lowercase Roman (i, ii, iii, iv...)
setup_page_number_footer(doc.sections[1], is_roman=True, start_num=1)

# Section 2: Main Report -> Arabic (1, 2, 3, 4...) restarting at 1
setup_page_number_footer(doc.sections[2], is_roman=False, start_num=1)

doc.save(DOC_PATH)
print("Footers and page numbering configured in document.")

# Step 2: Remove evenAndOddHeaders from settings.xml inside docx zip
with zipfile.ZipFile(DOC_PATH, 'r') as zin:
    file_map = {}
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == 'word/settings.xml':
            settings_str = data.decode('utf-8')
            settings_str = settings_str.replace('<w:evenAndOddHeaders/>', '')
            settings_str = settings_str.replace('<w:evenAndOddHeaders />', '')
            data = settings_str.encode('utf-8')
        file_map[item.filename] = data

with zipfile.ZipFile(DOC_PATH, 'w', zipfile.ZIP_DEFLATED) as zout:
    for filename, data in file_map.items():
        zout.writestr(filename, data)

print("Settings.xml cleaned up (evenAndOddHeaders removed).")
